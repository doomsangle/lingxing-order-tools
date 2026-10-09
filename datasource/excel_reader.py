"""
Excel 订单数据读取。

返回完整行数据（订单号 + ASIN + 金额 + 原始行），用于后续三重校验：
  1) 订单是否存在
  2) 订单ID 与 ASIN 是否匹配
  3) 订单金额是否相等（四舍五入保留2位小数）
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable


# 列标题（精确匹配优先，子串兜底）
ORDER_ID_EXACT = ["订单号"]
ORDER_ID_HINTS = ["订单号", "order id", "orderid", "order_id", "order number", "amazon order"]

ASIN_EXACT = ["sku或asin"]
ASIN_HINTS = ["sku或asin", "asin"]

AMOUNT_EXACT = ["下单实付金额"]
AMOUNT_HINTS = ["下单实付金额", "amount", "实付金额"]


@dataclass
class ExcelOrderRow:
    """Excel 中一行订单的完整数据。"""
    order_id: str
    asin: str = ""
    amount: float = 0.0
    raw_row: tuple = field(default_factory=tuple)
    sheet_name: str = ""
    row_idx: int = 0          # 在 sheet 中的行号（1-based，含表头）
    col_count: int = 0


def _find_col(header_row: Iterable, exact: list[str], hints: list[str]) -> int | None:
    """
    两阶段查找列索引：
      1) 精确匹配：单元格文本 strip 后 == 精确标题（不分大小写）
      2) 子串匹配：标题包含 hint 关键词
    """
    # 阶段1：精确匹配
    for idx, cell in enumerate(header_row):
        text = str(cell or "").strip()
        for name in exact:
            if text.lower() == name.lower():
                return idx
    # 阶段2：子串匹配
    for idx, cell in enumerate(header_row):
        text = str(cell or "").strip().lower()
        for hint in hints:
            if hint in text:
                return idx
    return None


def _parse_amount(raw) -> float:
    """安全解析金额，失败返回 0。"""
    if raw is None:
        return 0.0
    try:
        return float(raw)
    except (ValueError, TypeError):
        s = str(raw).strip().replace(",", "").replace("￥", "").replace("$", "")
        try:
            return float(s)
        except (ValueError, TypeError):
            return 0.0


def read_excel_orders(excel_path: str | Path) -> list[ExcelOrderRow]:
    """
    读取 xlsx / xls / csv 文件，返回每行订单的完整数据。

    CSV 视为单 sheet，sheet_name 用文件名（去扩展名）。
    """
    path = Path(excel_path)
    if not path.is_file():
        raise FileNotFoundError(f"文件不存在: {path}")

    ext = path.suffix.lower()
    if ext == ".csv":
        return _read_csv_orders(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx_orders(path)
    # .xls 用 openpyxl 读不了（老格式），用 xlrd
    if ext == ".xls":
        return _read_xls_orders(path)
    raise ValueError(f"不支持的文件类型: {ext}（仅支持 .xlsx / .xls / .csv）")


# ----------------------------------------------------------------------
def _read_csv_orders(path: Path) -> list[ExcelOrderRow]:
    import csv
    encoding = _detect_csv_encoding(path)
    orders: list[ExcelOrderRow] = []
    seen: set[str] = set()

    with open(path, "r", encoding=encoding, errors="replace", newline="") as f:
        reader = csv.reader(f)
        header = next(reader, None)
        if header is None:
            return orders
        oid_col = _find_col(header, ORDER_ID_EXACT, ORDER_ID_HINTS)
        if oid_col is None:
            raise ValueError("CSV 中未找到订单号列")
        asin_col = _find_col(header, ASIN_EXACT, ASIN_HINTS)
        amount_col = _find_col(header, AMOUNT_EXACT, AMOUNT_HINTS)

        for row_idx, row in enumerate(reader, 2):
            if oid_col >= len(row):
                continue
            raw = row[oid_col]
            if raw is None or raw == "":
                continue
            oid = re.sub(r"[\s\u200b-\u200f\ufeff]", "", str(raw).strip())
            if not oid or oid in seen:
                continue
            seen.add(oid)

            asin_val = str(row[asin_col]).strip() if asin_col is not None and asin_col < len(row) else ""
            amt = _parse_amount(row[amount_col]) if amount_col is not None and amount_col < len(row) else 0.0
            orders.append(ExcelOrderRow(
                order_id=oid, asin=asin_val, amount=amt,
                raw_row=tuple(row), sheet_name=path.stem, row_idx=row_idx,
                col_count=len(row),
            ))
    return orders


def _detect_csv_encoding(path: Path) -> str:
    """尝试 UTF-8 / GBK / UTF-8-SIG，常见中文 CSV 编码。"""
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            f.read(4096)
        return "utf-8-sig"
    except UnicodeDecodeError:
        pass
    try:
        with open(path, "r", encoding="gbk") as f:
            f.read(4096)
        return "gbk"
    except UnicodeDecodeError:
        pass
    return "utf-8"


# ----------------------------------------------------------------------
def _read_xlsx_orders(path: Path) -> list[ExcelOrderRow]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    orders: list[ExcelOrderRow] = []
    seen_ids: set[str] = set()

    for sheet in wb.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            continue

        oid_col = _find_col(header, ORDER_ID_EXACT, ORDER_ID_HINTS)
        if oid_col is None:
            continue  # 没有订单号列的 sheet 直接跳过

        asin_col = _find_col(header, ASIN_EXACT, ASIN_HINTS)
        amount_col = _find_col(header, AMOUNT_EXACT, AMOUNT_HINTS)

        for row_idx, row in enumerate(rows_iter, 2):  # 从第2行开始（1-based，第1行是表头）
            if oid_col >= len(row):
                continue
            raw_oid = row[oid_col]
            if raw_oid is None:
                continue
            oid = re.sub(r"[\s\u200b-\u200f\ufeff]", "", str(raw_oid).strip())
            if not oid or oid in seen_ids:
                continue
            seen_ids.add(oid)

            asin_val = ""
            if asin_col is not None and asin_col < len(row) and row[asin_col] is not None:
                asin_val = str(row[asin_col]).strip()

            amount_val = 0.0
            if amount_col is not None and amount_col < len(row) and row[amount_col] is not None:
                amount_val = _parse_amount(row[amount_col])

            orders.append(ExcelOrderRow(
                order_id=oid, asin=asin_val, amount=amount_val,
                raw_row=row, sheet_name=sheet.title, row_idx=row_idx,
                col_count=len(row),
            ))

    wb.close()
    return orders


# ----------------------------------------------------------------------
def _read_xls_orders(path: Path) -> list[ExcelOrderRow]:
    """老版 .xls 格式，用 xlrd 读（若未安装则给出明确提示）。"""
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError(
            f"读取 .xls 需要 xlrd 库，但当前环境未安装。\n"
            f"请把 .xls 另存为 .xlsx 或 .csv，再用本工具处理。"
        ) from exc

    wb = xlrd.open_workbook(str(path))
    orders: list[ExcelOrderRow] = []
    seen_ids: set[str] = set()

    for sheet in wb.sheets():
        if sheet.nrows == 0:
            continue
        header = sheet.row_values(0)
        oid_col = _find_col(header, ORDER_ID_EXACT, ORDER_ID_HINTS)
        if oid_col is None:
            continue
        asin_col = _find_col(header, ASIN_EXACT, ASIN_HINTS)
        amount_col = _find_col(header, AMOUNT_EXACT, AMOUNT_HINTS)

        for r in range(1, sheet.nrows):
            row = sheet.row_values(r)
            if oid_col >= len(row):
                continue
            raw_oid = row[oid_col]
            if raw_oid is None or raw_oid == "":
                continue
            oid = re.sub(r"[\s\u200b-\u200f\ufeff]", "", str(raw_oid).strip())
            if not oid or oid in seen_ids:
                continue
            seen_ids.add(oid)

            asin_val = str(row[asin_col]).strip() if asin_col is not None and asin_col < len(row) else ""
            amt = _parse_amount(row[amount_col]) if amount_col is not None and amount_col < len(row) else 0.0
            orders.append(ExcelOrderRow(
                order_id=oid, asin=asin_val, amount=amt,
                raw_row=tuple(row), sheet_name=sheet.name, row_idx=r + 1,
                col_count=len(row),
            ))
    return orders


# ---- 向后兼容 ----
def iter_order_ids_from_excel(excel_path: str | Path) -> list[str]:
    """旧接口：仅返回订单号列表。"""
    return [o.order_id for o in read_excel_orders(excel_path)]


def looks_like_order_id(value: str) -> bool:
    v = value.strip()
    return bool(re.compile(r"^\d[\d\-]{5,25}\d$").match(v))
